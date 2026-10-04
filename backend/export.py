"""
Génération de rapports Excel complets avec openpyxl.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from backend.analysis import get_descriptive_stats, get_simple_regression, get_multiple_regression, get_logistic_regression

def generate_excel_report(df, output_path):
    """Génère un fichier Excel avec 6 feuilles d'analyse."""
    wb = Workbook()
    
    # Feuille 1: Données brutes
    ws1 = wb.active
    ws1.title = "Données Brutes"
    for r in dataframe_to_rows(df, index=False, header=True):
        ws1.append(r)
        
    # Feuille 2: Statistiques descriptives
    ws2 = wb.create_sheet("Statistiques Descriptives")
    desc_stats = get_descriptive_stats(df)
    ws2.append(["Variable", "Moyenne", "Médiane", "Min", "Max", "Écart-type", "CV"])
    for var, stats in desc_stats['stats'].items():
        ws2.append([var, stats['moyenne'], stats['mediane'], stats['min'], stats['max'], stats['ecart_type'], stats['cv']])
        
    # Feuille 3: Régression Simple
    ws3 = wb.create_sheet("Régression Simple")
    reg_simple = get_simple_regression(df)
    ws3.append(["Paramètre", "Valeur"])
    ws3.append(["Beta_0 (Intercept)", reg_simple['beta_0']])
    ws3.append(["Beta_1 (Revenu)", reg_simple['beta_1']])
    ws3.append(["R²", reg_simple['r2']])
    ws3.append(["Prédiction pour 750 000 FCFA", reg_simple['prediction_result']])
    
    # Feuille 4: Régression Multiple
    ws4 = wb.create_sheet("Régression Multiple")
    reg_multiple = get_multiple_regression(df)
    ws4.append(["Paramètre", "Valeur"])
    for coeff_name, values in reg_multiple['coefficients'].items():
        ws4.append([f"Coeff {coeff_name} (Standardisé)", values['standardized']])
    ws4.append(["R²", reg_multiple['r2']])
    ws4.append(["RMSE", reg_multiple['rmse']])
    
    # Feuille 5: Régression Logistique
    ws5 = wb.create_sheet("Régression Logistique")
    reg_log = get_logistic_regression(df)
    ws5.append(["Métrique", "Valeur"])
    for metric, value in reg_log['metrics'].items():
        ws5.append([metric.upper(), value])
    ws5.append(["Matrice de Confusion (TN, FP / FN, TP)", str(reg_log['confusion_matrix'].tolist())])
    
    # Feuille 6: Tableau de bord KPI
    ws6 = wb.create_sheet("Tableau de Bord KPI")
    ws6.append(["Indicateur", "Valeur"])
    ws6.append(["Nombre total de clients", len(df)])
    ws6.append(["Montant total des crédits", df['Montant_Credit'].sum()])
    ws6.append(["Revenu moyen", df['Revenu_Mensuel'].mean()])
    ws6.append(["Taux de défaut (%)", desc_stats['default_rate']])
    
    wb.save(output_path)